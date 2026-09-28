"""사용자 Q&A 라이브 점검 JSON을 Word 보고서로 렌더링한다."""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt


def _status_label(status: str) -> str:
    return {"PASS": "PASS", "BLOCKED_DATA": "SKIP (원천 부재)",
            "PARTIAL": "FAIL (부분 응답)"}.get(status, "FAIL")


def _set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(8)


def _debug_text(row: dict) -> str:
    """SSE DEBUG 이벤트를 Word에서 바로 확인할 수 있는 짧은 진단으로 만든다."""
    messages: list[str] = []
    for event in row.get("debug", []) or []:
        for result in event.get("action_results", []) or []:
            if result.get("status") != "success":
                messages.append(
                    f"{result.get('requirement_id', 'unknown')}/"
                    f"{result.get('action_id', 'unknown')}: "
                    f"{result.get('failure_reason') or result.get('status') or 'unknown'}"
                )
        for warning in event.get("warnings", []) or []:
            messages.append(f"경고: {warning}")
    return "\n".join(dict.fromkeys(messages)) or "-"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()

    payload = json.loads(args.input_json.read_text(encoding="utf-8"))
    results = payload["results"]
    counts = Counter(row["status"] for row in results)

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.7)
    normal = document.styles["Normal"]
    normal.font.size = Pt(9)

    document.add_heading("사용자 제공 Q&A 쌍 라이브 재점검", level=0)
    document.add_paragraph(
        f"실행 시각: {payload.get('generated_at')} (Asia/Seoul)\n"
        f"대상: {payload.get('base_url')}/pubchat\n"
        f"질문 수: {len(results)}건\n"
        f"집계: " + ", ".join(f"{key} {value}" for key, value in sorted(counts.items()))
    )
    if scope_note := payload.get("scope_note"):
        document.add_paragraph(f"검증 범위: {scope_note}")
    document.add_paragraph(
        "PASS는 SSE 종료 계약과 질문별 기대 표지가 확인된 응답입니다. "
        "SKIP은 source_unavailable 안전 종료이며 구현 실패로 집계하지 않습니다. "
        "표·차트는 존재 여부만 기록하고 시각적 품질은 이번 판정에서 제외합니다. "
        "DEBUG가 활성화된 실행은 SSE 디버그 이벤트의 실패 Action·경고를 마지막 열에 기록합니다."
    )

    document.add_heading("질문별 결과 및 실제 출력", level=1)
    # 요약과 실제 답변을 분리하지 않는다. Word를 열자마자 각 질문의 판정과
    # 답변을 같은 행에서 읽을 수 있게 해, 별도 페이지의 본문을 놓치는 일을 막는다.
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("ID", "판정", "질문", "실제 출력", "Action / Source", "비고", "디버그 진단")):
        _set_cell_text(cell, text)
    for row in results:
        cells = table.add_row().cells
        source = ", ".join(row.get("actions", []) + row.get("sources", [])) or "-"
        notes = "; ".join(row.get("notes", [])) or "-"
        answer = row.get("answer") or "(응답 없음)"
        debug = _debug_text(row)
        question = row["question"]
        if row.get("original_question"):
            question += f"\n원문: {row['original_question']}\n날짜 보정: {row.get('qa_adjustment') or '-'}"
        for cell, text in zip(cells, (row["id"], _status_label(row["status"]), question, answer, source, notes, debug)):
            _set_cell_text(cell, text)

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)


if __name__ == "__main__":
    main()
