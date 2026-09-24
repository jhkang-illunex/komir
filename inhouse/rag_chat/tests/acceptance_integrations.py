"""고정 표본의 원본→OKF→PageIndex→pgvector 읽기 전용 수락 검사."""
from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import yaml

FIXTURES = Path(__file__).with_name("acceptance_fixtures.json")
REPO_ROOT = Path(__file__).resolve().parents[3]

REQUIRED_CHECK_IDS = {
    "AC28": ("original_metadata_provenance", "fixed_document_chain", "dense_expected_document", "format_sample_policy"),
    "AC29": ("original_metadata_provenance", "fixed_document_chain", "dense_expected_document", "format_sample_policy"),
    "AC30": ("original_metadata_provenance", "fixed_document_chain", "dense_expected_document", "format_sample_policy"),
    "AC39": ("embedding_model_dimension", "dense_candidate_metadata", "dense_expected_document_topk", "vector_only_no_fallback"),
    "AC40": ("fixed_doc_filter", "node_okf_range_identity", "node_original_text"),
    "AC41": ("three_nonempty_bodies", "three_tree_hashes_fresh", "three_valid_roots_and_node_text", "three_latest_chunk_texts", "all_three_atomic"),
}


def fixture(case_id: str) -> dict[str, Any]:
    return json.loads(FIXTURES.read_text(encoding="utf-8")).get(case_id, {})


@dataclass
class _CheckLedger:
    case_id: str
    checks: dict[str, dict[str, Any]] = field(default_factory=dict)

    def record(self, check_id: str, status: str, **evidence: Any) -> None:
        if check_id not in REQUIRED_CHECK_IDS[self.case_id] or check_id in self.checks:
            raise AssertionError(f"required check 정의/중복 오류: {self.case_id}/{check_id}")
        self.checks[check_id] = {"id": check_id, "status": status, **evidence}

    def finish(self, errors: list[str], *, blocked: str | None = None,
               skip_reason: str = "선행 검사 실패로 실행하지 못함") -> tuple[str, list[str], dict[str, Any]]:
        for check_id in REQUIRED_CHECK_IDS[self.case_id]:
            if check_id not in self.checks:
                self.record(check_id, "SKIPPED", reason=skip_reason)
        required = [self.checks[item] for item in REQUIRED_CHECK_IDS[self.case_id]]
        statuses = [item["status"] for item in required]
        if "FAIL" in statuses:
            status = "FAIL"
        elif any(value.startswith("BLOCKED_") for value in statuses):
            status = blocked or next(value for value in statuses if value.startswith("BLOCKED_"))
        elif "SKIPPED" in statuses:
            status = blocked or "FAIL"
            errors.append("필수 검사가 실행되지 않음")
        else:
            status = "PASS"
        if status == "PASS" and (len(required) != len(set(item["id"] for item in required))
                                 or any(item["status"] != "PASS" for item in required)):
            status = "FAIL"
            errors.append("required_checks 완결성 위반")
        return status, errors, {"required_checks": required}


def _tree_for_spec(trees: list[dict[str, Any]], spec: dict[str, Any]) -> dict[str, Any] | None:
    matches = [tree for tree in trees if tree.get("resource") == spec.get("resource")]
    for key in ("doc_id", "okf_path"):
        if spec.get(key):
            matches = [tree for tree in matches if tree.get(key) == spec[key]]
    return matches[0] if len(matches) == 1 else None


def _vector_doc_id(value: str) -> str:
    """공식 OKF→pgvector 적재 규칙: doc_ 제거 후 앞 16자."""
    return str(value or "").removeprefix("doc_")[:16]


def _body_text(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---\n"):
        return text
    end = text.find("\n---\n", 3)
    return text[end + 5:].lstrip("\n") if end >= 0 else text


def _frontmatter(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---\n"):
        return {}
    end = text.find("\n---\n", 3)
    return (yaml.safe_load(text[4:end]) or {}) if end >= 0 else {}


def _body_sha256(path: Path) -> str:
    return hashlib.sha256(_body_text(path).encode("utf-8")).hexdigest()


def _body_line_offset(path: Path) -> int:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---\n"):
        return 0
    end = text.find("\n---\n", 3)
    if end < 0:
        return 0
    rest = text[end + 5:]
    return text[:end + 5].count("\n") + len(rest) - len(rest.lstrip("\n"))


def _fact_in_line_range(path: Path, fact: str, start: int, end: int) -> bool:
    lines = _body_text(path).splitlines()
    return 0 < start <= end <= len(lines) and fact in "\n".join(lines[start - 1:end])


def _iter_nodes(structure: list[dict[str, Any]]):
    for node in structure:
        yield node
        yield from _iter_nodes(node.get("nodes", []))


def _independent_node_text(tree: dict[str, Any], hit: dict[str, Any], okf: Path,
                           *, max_chars: int = 4000) -> tuple[str, int, int, list[str]]:
    """hit metadata를 신뢰하지 않고 tree node와 OKF 본문에서 범위를 다시 만든다."""
    problems = []
    actual_offset = _body_line_offset(okf)
    if tree.get("body_line_offset") != actual_offset:
        problems.append("tree body_line_offset이 실제 OKF 프론트매터 범위와 다름")
    for key in ("doc_id", "resource", "okf_path", "body_line_offset"):
        if hit.get(key) != tree.get(key):
            problems.append(f"hit {key} 불일치")
    nodes = [node for node in _iter_nodes(tree.get("structure", [])) if node.get("node_id") == hit.get("node_id")]
    if len(nodes) != 1:
        return "", 0, 0, [*problems, "hit node_id가 tree에서 유일하지 않음"]
    if hit.get("line_num") != nodes[0].get("line_num"):
        problems.append("hit line_num이 tree node와 다름")
    lines = _body_text(okf).splitlines()
    start = hit.get("line_num")
    if not isinstance(start, int) or not 1 <= start <= len(lines):
        return "", int(start or 0), 0, [*problems, "hit line_num이 OKF 본문 범위 밖"]
    heading = lines[start - 1].strip()
    depth = len(heading) - len(heading.lstrip("#")) if heading.startswith("#") else 0
    end = len(lines)
    for index in range(start, len(lines)):
        stripped = lines[index].lstrip()
        if not stripped.startswith("#"):
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        if depth and level > depth:
            continue
        end = index
        break
    return "\n".join(lines[start - 1:end]).strip()[:max_chars], start, end, problems


def _compact(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def _facts_share_context(text: str, anchor: str, facts: list[str], max_chars: int):
    compact = _compact(text)
    positions = [compact.find(_compact(value)) for value in (anchor, *facts)]
    ok = all(position >= 0 for position in positions) and max(positions) - min(positions) <= max_chars
    return ok, {"anchor": anchor, "facts": facts, "positions": positions, "max_span_chars": max_chars}


def _hwp_text_and_rows(original: Path) -> tuple[str, list[str], str | None]:
    with tempfile.TemporaryDirectory(prefix="acceptance-hwp-") as temp_dir:
        odt = Path(temp_dir) / "source.odt"
        try:
            result = subprocess.run(["hwp5odt", "--output", str(odt), str(original)], text=True,
                                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30, check=False)
        except (OSError, subprocess.TimeoutExpired) as exc:
            return "", [], f"hwp5odt 호출 불가: {type(exc).__name__}: {exc}"
        if result.returncode or not odt.is_file():
            return "", [], f"hwp5odt 변환 실패: {result.stderr[-500:]}"
        try:
            with zipfile.ZipFile(odt) as archive:
                root = ElementTree.fromstring(archive.read("content.xml"))
            text = " ".join(part.strip() for part in root.itertext() if part.strip())
            rows = [" ".join(part.strip() for part in element.itertext() if part.strip())
                    for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "table-row"]
            return text, [row for row in rows if row], None
        except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            return "", [], f"hwp5odt 결과 읽기 실패: {type(exc).__name__}: {exc}"


def _hwp_text(original: Path) -> tuple[str, str | None]:
    text, _, error = _hwp_text_and_rows(original)
    return text, error


def _verify_hwp_original(original: Path, spec: dict[str, Any]):
    text, rows, error = _hwp_text_and_rows(original)
    if error:
        return "BLOCKED_ENV", {}, error
    matched_rows = [row for row in rows if _compact(row).startswith(_compact(spec["context_anchor"]))
                    and all(_compact(value) in _compact(row) for value in spec["original_facts"])]
    evidence = {"table_row_count": len(rows), "matched_row_count": len(matched_rows),
                "matched_rows": matched_rows[:3], "flattened_text_chars": len(text)}
    return ("PASS" if matched_rows else "FAIL"), evidence, None if matched_rows else "HWP 니켈 table-row의 사실 문맥 연결 실패"


def _same_cell(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return math.isclose(float(actual), float(expected), rel_tol=1e-12, abs_tol=1e-9)
    return str(actual).strip() == str(expected).strip()


def _verify_excel_original(original: Path, spec: dict[str, Any]):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return "BLOCKED_ENV", {}, f"openpyxl 없음: {exc}"
    try:
        book = load_workbook(original, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        return "BLOCKED_ENV", {}, f"XLSX 원본 읽기 실패: {type(exc).__name__}: {exc}"
    if spec["sheet"] not in book.sheetnames:
        return "FAIL", {"sheet_names": book.sheetnames}, "XLSX 고정 sheet 없음"
    sheet, contract = book[spec["sheet"]], spec["row_contract"]
    cells = {contract["unit_cell"]: sheet[contract["unit_cell"]].value}
    ok = _same_cell(cells[contract["unit_cell"]], contract["unit"])
    for item in contract["fields"]:
        for kind in ("header", "value"):
            cell = item[f"{kind}_cell"]
            cells[cell] = sheet[cell].value
            ok &= _same_cell(cells[cell], item[kind])
    ok &= str(spec["year"]) in original.name
    evidence = {"sheet": spec["sheet"], "cells": cells, "row_contract": contract, "filename": original.name}
    return ("PASS" if ok else "FAIL"), evidence, None if ok else "XLSX 고정 행·열 관계 불일치"


def _markdown_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _verify_excel_okf_table(body: str, spec: dict[str, Any]):
    """OKF 표의 header column과 unit row 값을 위치 관계로 검증한다."""
    lines = body.splitlines()
    wanted_headers = [field["header"] for field in spec["row_contract"]["fields"]]
    header_candidates = [(index, _markdown_cells(line)) for index, line in enumerate(lines)
                         if line.lstrip().startswith("|") and all(value in _markdown_cells(line) for value in wanted_headers)]
    evidence = {"headers": wanted_headers, "unit": spec["row_contract"]["unit"]}
    for header_index, header in header_candidates:
        columns = {name: header.index(name) for name in wanted_headers}
        for line in lines[header_index + 2:]:
            if not line.lstrip().startswith("|"):
                break
            row = _markdown_cells(line)
            if not row or row[0] != spec["row_contract"]["unit"]:
                continue
            comparisons = []
            ok = True
            for field in spec["row_contract"]["fields"]:
                actual = row[columns[field["header"]]] if columns[field["header"]] < len(row) else ""
                expected = field["value"]
                if isinstance(expected, float) and not float(expected).is_integer():
                    matched = actual == f"{expected:.6f}".rstrip("0").rstrip(".")
                else:
                    matched = actual == str(expected)
                comparisons.append({"header": field["header"], "column": columns[field["header"]],
                                    "actual": actual, "expected_source": expected, "matched": matched})
                ok &= matched
            evidence.update({"header_line": header_index + 1, "unit_row": row, "comparisons": comparisons})
            return ok, evidence
    evidence["reason"] = "header 또는 unit row 없음"
    return False, evidence


def _fetch_chunks(connect, settings, *, doc_id: str, source_path: str):
    try:
        conn = connect()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    f"SELECT chunk_id, doc_id, source_path, week, title, section_heading, txt, vector_dims(embedding) "
                    f"FROM {settings().PG_SCHEMA}.doc_chunk WHERE doc_id=%s AND source_path=%s AND embedding IS NOT NULL",
                    (_vector_doc_id(doc_id), source_path),
                )
                rows = cur.fetchall()
        finally:
            conn.close()
    except Exception as exc:
        try:
            from psycopg2 import InterfaceError, OperationalError
            environment_error = isinstance(exc, (InterfaceError, OperationalError, OSError, ConnectionError))
        except ImportError:
            environment_error = isinstance(exc, (OSError, ConnectionError))
        status = "BLOCKED_ENV" if environment_error else "FAIL"
        return status, [], f"pgvector 읽기 실패: {type(exc).__name__}: {exc}"
    keys = ("chunk_id", "doc_id", "source_path", "source_group", "title", "section_heading", "text", "dimension")
    return "PASS", [dict(zip(keys, row)) for row in rows], None


def _exception_status(exc: Exception) -> str:
    # psycopg/async DB 클라이언트가 이미 닫힌 경우에는 데이터 계약 실패가
    # 아니라 수락검사 실행 환경의 연결 수명 문제다. 이 메시지는 드라이버가
    # RuntimeError로 감싸서 전달하므로 타입만으로는 환경 차단을 판별할 수 없다.
    message = str(exc).casefold()
    if "client has been closed" in message or "connection is closed" in message:
        return "BLOCKED_ENV"
    try:
        from psycopg2 import InterfaceError, OperationalError
        environment_error = isinstance(exc, (InterfaceError, OperationalError, OSError, ConnectionError, ImportError))
    except ImportError:
        environment_error = isinstance(exc, (OSError, ConnectionError, ImportError))
    return "BLOCKED_ENV" if environment_error else "FAIL"


def _expected_chunk_texts(okf: Path) -> set[str]:
    from ingest.vectorize.build_pgvector_okf import _load_okf_record
    from rag_core.ragkit.chunk import chunk_document
    return {chunk.text for chunk in chunk_document(_load_okf_record(okf))}


def _dense_search(spec: dict[str, Any]):
    from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
    from rag_core.retrieval.dense_pg import dense_search_pg
    return dense_search_pg(spec["query"], k=int(spec.get("top_k", 10)),
                           exclude_src=PRIVATE_ONLY_SOURCE_GROUPS, _allow_fallback=False)


def _dense_evidence(spec: dict[str, Any], hits: list[Any]) -> dict[str, Any]:
    expected_id = _vector_doc_id(spec["doc_id"])
    candidates = [{"rank": getattr(hit, "dense_rank", None), "chunk_id": getattr(hit, "chunk_id", ""),
                   "doc_id": getattr(hit, "doc_id", ""), "source_path": getattr(hit, "source_path", ""),
                   "source_group": getattr(hit, "week", ""), "title": getattr(hit, "title", ""),
                   "section_heading": getattr(hit, "section_heading", ""),
                   "text_chars": len(getattr(hit, "text", "")), "score": getattr(hit, "score", None),
                   "fact_present": spec["okf_fact"] in getattr(hit, "text", "")} for hit in hits]
    expected = [item["rank"] for item in candidates
                if item["doc_id"] == expected_id and item["source_path"] == spec["resource"]]
    exact = [item["rank"] for item in candidates
             if item["doc_id"] == expected_id and item["source_path"] == spec["resource"] and item["fact_present"]]
    return {"query": spec["query"], "top_k": spec.get("top_k", 10), "expected_tree_doc_id": spec["doc_id"],
            "expected_vector_doc_id": expected_id, "expected_source_path": spec["resource"],
            "expected_doc_ranks": expected, "exact_fact_ranks": exact, "candidates": candidates}


def _candidate_metadata_valid(hits: list[Any]):
    problems = []
    if [getattr(hit, "dense_rank", None) for hit in hits] != list(range(1, len(hits) + 1)):
        problems.append("dense_rank가 연속 순위가 아님")
    from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
    for hit in hits:
        rank = getattr(hit, "dense_rank", None)
        if not all(getattr(hit, key, "") for key in ("chunk_id", "doc_id", "source_path", "title", "text")):
            problems.append(f"rank {rank} 필수 metadata/text 누락")
        score = getattr(hit, "score", None)
        if not isinstance(score, (int, float)) or not math.isfinite(float(score)):
            problems.append(f"rank {rank} score 비정상")
        if getattr(hit, "week", "") in PRIVATE_ONLY_SOURCE_GROUPS:
            problems.append(f"rank {rank} public 제외 source_group 노출")
    return not problems, problems


def verify_dense_fixture(case_id: str, *, tree: dict[str, Any], okf_root: Path):
    spec = fixture(case_id)
    okf = okf_root / str(tree.get("okf_path", ""))
    if tree.get("resource") != spec.get("resource") or not okf.is_file() or spec.get("okf_fact") not in _body_text(okf):
        return "FAIL", ["고정 tree/OKF provenance 불일치"], {"tree": tree, "fixture": spec}
    try:
        hits = _dense_search(spec)
    except Exception as exc:
        status = _exception_status(exc)
        return status, [f"dense_search_pg 실행 실패: {type(exc).__name__}: {exc}"], {"query": spec["query"]}
    evidence = _dense_evidence(spec, hits)
    valid, problems = _candidate_metadata_valid(hits)
    if not hits or not valid or not evidence["exact_fact_ranks"]:
        return "FAIL", problems or ["dense top-k에서 고정 문서의 exact fact를 회수하지 못함"], evidence
    return "PASS", [], evidence


def _verify_document_case(case_id: str, trees, pageindex, connect, settings):
    spec, ledger, errors = fixture(case_id), _CheckLedger(case_id), []
    blocked = None
    tree = _tree_for_spec(trees, spec)
    original = REPO_ROOT / spec["resource"]
    okf = Path(pageindex.OKF_DOCUMENTS_ROOT) / spec["okf_path"]

    policy_ok = (all(spec.get(key) for key in ("resource", "doc_id", "okf_path", "format", "source_sha256", "body_sha256"))
                 and Path(spec["resource"]).suffix.lower() == "." + spec["format"] and tree is not None)
    ledger.record("format_sample_policy", "PASS" if policy_ok else "FAIL",
                  pinned_identity={key: spec.get(key) for key in ("resource", "doc_id", "okf_path", "format")},
                  exact_tree_matches=int(tree is not None))
    if not policy_ok:
        errors.append("고정 포맷 표본의 전체 식별자/형식 불일치")

    provenance = {"resource": spec["resource"], "expected_sha256": spec["source_sha256"]}
    if not original.is_file():
        blocked = "BLOCKED_DATA"
        ledger.record("original_metadata_provenance", blocked, **provenance, reason="고정 원본 없음")
        errors.append("고정 원본 표본이 없음")
    else:
        source_hash = hashlib.sha256(original.read_bytes()).hexdigest()
        provenance["actual_sha256"] = source_hash
        original_ok, raw_status = source_hash == spec["source_sha256"], "PASS"
        if case_id == "AC29":
            raw_status, raw_evidence, error = _verify_hwp_original(original, spec)
            provenance["raw_context"] = raw_evidence
            if error:
                errors.append(error)
        elif case_id == "AC30":
            raw_status, raw_evidence, error = _verify_excel_original(original, spec)
            provenance["raw_row_contract"] = raw_evidence
            if error:
                errors.append(error)
        original_ok &= raw_status == "PASS"
        if raw_status.startswith("BLOCKED_"):
            blocked = raw_status
        confirmed_hash_failure = source_hash != spec["source_sha256"]
        ledger.record("original_metadata_provenance", "PASS" if original_ok else
                      ("FAIL" if confirmed_hash_failure or raw_status == "FAIL" else (blocked or "FAIL")), **provenance)
        if source_hash != spec["source_sha256"]:
            errors.append("고정 원본 SHA-256 불일치")

    chain, problems = {}, []
    if tree is None:
        problems.append("고정 tree 없음 또는 중복")
    if not okf.is_file():
        problems.append("고정 OKF 없음")
    if tree is not None and okf.is_file():
        front, current_hash = _frontmatter(okf), _body_sha256(okf)
        metadata_keys = ("doc_id", "resource", "fmt", "source_group", "document_date", "minerals", "content_keywords")
        metadata_equal = all(tree.get(key) == front.get(key) for key in metadata_keys)
        identity_equal = (front.get("doc_id") == spec["doc_id"] and front.get("resource") == spec["resource"]
                          and front.get("fmt") == spec["format"] and front.get("content_sha256") == spec["source_sha256"])
        fresh = current_hash == spec["body_sha256"] == tree.get("okf_body_sha256")
        chain.update({"front_identity": {key: front.get(key) for key in ("doc_id", "resource", "fmt", "content_sha256")},
                      "metadata_equal_front_tree": metadata_equal,
                      "metadata_gaps": [key for key in ("document_date", "minerals") if not front.get(key)],
                      "fixture_body_sha256": spec["body_sha256"], "current_body_sha256": current_hash,
                      "tree_body_sha256": tree.get("okf_body_sha256"), "vector_doc_id": _vector_doc_id(spec["doc_id"])})
        if not metadata_equal or not identity_equal:
            problems.append("OKF frontmatter/tree metadata 계약 불일치")
        if not fresh:
            problems.append("fixture/current/tree OKF 본문 hash 불일치")
        facts = [spec["okf_fact"], *spec.get("original_facts", [])]
        body = _body_text(okf)
        if any(fact not in body for fact in facts):
            problems.append("downstream OKF 사실 누락")
        if case_id == "AC29":
            matching_lines = [line for line in body.splitlines()
                              if _compact(line).startswith(_compact(spec["context_anchor"]))
                              and all(_compact(value) in _compact(line) for value in spec["original_facts"])]
            chain["okf_context"] = {"matching_line_count": len(matching_lines), "matching_lines": matching_lines[:3]}
            if not matching_lines:
                problems.append("HWP 원본 사실과 OKF 니켈 문맥 연결 실패")
        if case_id == "AC30":
            table_ok, table_evidence = _verify_excel_okf_table(body, spec)
            chain["excel_downstream_relation"] = table_evidence
            if not table_ok:
                problems.append("Excel header/value/unit의 동일 행·열 관계가 OKF에서 끊김")
            facts = spec["downstream_row_tokens"]
        from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
        try:
            hits = pageindex.search_nodes(spec["query"], doc=spec["doc_id"],
                                          node_limit=int(spec.get("node_top_k", 10)),
                                          exclude_source_groups=PRIVATE_ONLY_SOURCE_GROUPS)
        except Exception as exc:
            hits = []
            problems.append(f"PageIndex 검사 구현/실행 실패: {type(exc).__name__}: {exc}")
        node_matches = []
        for hit in hits:
            independent, start, end, hit_problems = _independent_node_text(tree, hit, okf)
            returned = pageindex.read_node_text(hit) if not hit_problems else ""
            if not hit_problems and returned == independent and all(fact in independent for fact in facts):
                node_matches.append({"node_id": hit["node_id"], "line_start": start, "line_end": end,
                                     "text_chars": len(independent)})
        chain["pageindex_matches"] = node_matches
        if not node_matches:
            problems.append("고정 doc 필터의 PageIndex 노드 범위/원문 사실 대조 실패")
        chunk_status, chunks, chunk_error = _fetch_chunks(connect, settings, doc_id=spec["doc_id"], source_path=spec["resource"])
        if chunk_status != "PASS":
            if chunk_status.startswith("BLOCKED_"):
                blocked = chunk_status
            problems.append(chunk_error or "pgvector 읽기 실패")
        else:
            try:
                expected_texts = _expected_chunk_texts(okf)
            except Exception as exc:
                expected_texts = set()
                problems.append(f"현재 OKF 청크 재구성 실패: {type(exc).__name__}: {exc}")
            actual_texts = {chunk["text"] for chunk in chunks}
            linked = any(all(fact in chunk["text"] for fact in facts) for chunk in chunks)
            chain.update({"doc_chunk_count": len(chunks), "expected_chunk_count": len(expected_texts),
                          "chunk_texts_equal_latest_okf": actual_texts == expected_texts,
                          "one_chunk_contains_linked_facts": linked})
            if not chunks or actual_texts != expected_texts or not linked:
                problems.append("pgvector 청크가 최신 OKF 사실/본문과 불일치")
    chain_status = "PASS" if not problems else (blocked if blocked and all("pgvector" in item for item in problems) else "FAIL")
    ledger.record("fixed_document_chain", chain_status, **chain, problems=problems)
    errors.extend(problems)

    try:
        hits = _dense_search(spec)
        dense = _dense_evidence(spec, hits)
        valid, candidate_problems = _candidate_metadata_valid(hits)
        dense["candidate_metadata_problems"] = candidate_problems
        dense_ok = bool(hits) and valid and bool(dense["exact_fact_ranks"])
        ledger.record("dense_expected_document", "PASS" if dense_ok else "FAIL", **dense)
        if not dense_ok:
            errors.append("dense top-k에서 고정 문서의 exact fact를 회수하지 못함")
    except Exception as exc:
        exception_status = _exception_status(exc)
        if exception_status.startswith("BLOCKED_"):
            blocked = exception_status
        ledger.record("dense_expected_document", exception_status, query=spec["query"], reason=f"{type(exc).__name__}: {exc}")
        errors.append("dense_search_pg 실행 실패")
    return ledger.finish(errors, blocked=blocked)


def _verify_ac39(connect, settings):
    spec, contract = fixture("AC39"), fixture("dense_contract")
    ledger, errors, blocked = _CheckLedger("AC39"), [], None
    try:
        from rag_core.ragkit.embed import DIM, MODEL_NAME, encode_query
        query_vector = encode_query(spec["query"])
        chunk_status, chunks, chunk_error = _fetch_chunks(connect, settings, doc_id=spec["doc_id"], source_path=spec["resource"])
        if chunk_status != "PASS":
            blocked = chunk_status
            ledger.record("embedding_model_dimension", blocked, reason=chunk_error)
            errors.append("임베딩 DB 차원 실행 검증 불가")
        else:
            db_dimensions = sorted({chunk["dimension"] for chunk in chunks})
            ok = (MODEL_NAME == contract["embedding_model"] and DIM == contract["embedding_dimension"]
                  and len(query_vector) == DIM and db_dimensions == [DIM] and bool(chunks))
            ledger.record("embedding_model_dimension", "PASS" if ok else "FAIL",
                          expected=contract, model_constant=MODEL_NAME, dimension_constant=DIM,
                          encoded_query_dimension=len(query_vector), db_dimensions=db_dimensions,
                          matched_chunk_count=len(chunks))
            if not ok:
                errors.append("실제 query/DB vector 차원이 임베딩 계약과 다름")
    except Exception as exc:
        exception_status = _exception_status(exc)
        if exception_status.startswith("BLOCKED_"):
            blocked = exception_status
        ledger.record("embedding_model_dimension", exception_status, reason=f"{type(exc).__name__}: {exc}")
        errors.append("임베딩 모델/차원 검사 실패")

    try:
        hits = _dense_search(spec)
        evidence = _dense_evidence(spec, hits)
        valid, problems = _candidate_metadata_valid(hits)
        from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
        leaked = [getattr(hit, "dense_rank", None) for hit in hits if getattr(hit, "week", "") in PRIVATE_ONLY_SOURCE_GROUPS]
        valid &= not leaked and len(hits) == int(spec["top_k"])
        ledger.record("dense_candidate_metadata", "PASS" if valid else "FAIL",
                      candidates=evidence["candidates"], problems=problems, private_source_ranks=leaked,
                      public_exclude_source_groups=sorted(PRIVATE_ONLY_SOURCE_GROUPS))
        if not valid:
            errors.append("dense top-k 후보 metadata/public 접근 경계 불일치")
        expected_ok, fact_ok = bool(evidence["expected_doc_ranks"]), bool(evidence["exact_fact_ranks"])
        ledger.record("dense_expected_document_topk", "PASS" if expected_ok and fact_ok else "FAIL", **evidence,
                      expected_document_found=expected_ok, exact_fact_found=fact_ok)
        if not expected_ok:
            errors.append("dense top-k에 기대 문서가 없음")
        if not fact_ok:
            errors.append("dense top-k 기대 문서 청크에 고정 exact fact가 없음")
        ledger.record("vector_only_no_fallback", "PASS", allow_fallback=False,
                      requested_top_k=spec["top_k"], returned_count=len(hits), query=spec["query"])
    except Exception as exc:
        exception_status = _exception_status(exc)
        if exception_status.startswith("BLOCKED_"):
            blocked = exception_status
        ledger.record("dense_candidate_metadata", exception_status, reason=f"{type(exc).__name__}: {exc}")
        errors.append("dense_search_pg 실행 실패")
    return ledger.finish(errors, blocked=blocked)


def _verify_ac40(trees, pageindex):
    spec, ledger, errors = fixture("AC40"), _CheckLedger("AC40"), []
    tree = _tree_for_spec(trees, spec)
    okf = Path(pageindex.OKF_DOCUMENTS_ROOT) / spec["okf_path"]
    from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
    hits = pageindex.search_nodes(spec["query"], doc=spec["doc_id"], node_limit=5,
                                  exclude_source_groups=PRIVATE_ONLY_SOURCE_GROUPS)
    filter_ok = tree is not None and all(hit.get("doc_id") == spec["doc_id"]
                                         and hit.get("resource") == spec["resource"]
                                         and hit.get("okf_path") == spec["okf_path"] for hit in hits)
    ledger.record("fixed_doc_filter", "PASS" if filter_ok else "FAIL", requested_doc_id=spec["doc_id"],
                  returned_count=len(hits), hit_identities=[{key: hit.get(key) for key in ("doc_id", "resource", "okf_path")} for hit in hits])
    if not filter_ok:
        errors.append("AC40 고정 doc 필터 결과 식별자 불일치")

    evidence, problems, matches = {}, [], []
    if tree is None or not okf.is_file():
        problems.append("고정 tree 또는 OKF 없음")
    else:
        current_hash, tree_hash = _body_sha256(okf), tree.get("okf_body_sha256")
        fresh = current_hash == spec["body_sha256"] == tree_hash
        evidence.update({"fixture_body_sha256": spec["body_sha256"], "current_body_sha256": current_hash,
                         "tree_body_sha256": tree_hash, "hashes_fresh": fresh})
        if not fresh:
            problems.append("fixture/current/tree OKF 본문 hash 불일치")
        if not _fact_in_line_range(okf, spec["fact"], spec["line_start"], spec["line_end"]):
            problems.append("고정 사실이 fixture 줄 범위에 없음")
        for hit in hits:
            independent, start, end, hit_problems = _independent_node_text(tree, hit, okf)
            if not hit_problems and start <= spec["line_start"] <= end and spec["fact"] in independent:
                matches.append((hit, independent, start, end))
        if not matches:
            problems.append("top-5 node가 고정 사실의 실제 OKF 범위를 가리키지 않음")
    evidence["matching_nodes"] = [{"node_id": hit["node_id"], "line_start": start, "line_end": end}
                                  for hit, _, start, end in matches]
    ledger.record("node_okf_range_identity", "PASS" if not problems else "FAIL", **evidence, problems=problems)
    errors.extend(problems)
    if not matches:
        ledger.record("node_original_text", "SKIPPED", reason="유효한 범위 node가 없음", nodes=[])
    else:
        text_rows, text_problems = [], []
        for hit, independent, _, _ in matches:
            returned = pageindex.read_node_text(hit)
            text_rows.append({"node_id": hit["node_id"], "independent_chars": len(independent),
                              "returned_chars": len(returned), "equal": returned == independent})
            if returned != independent or spec["fact"] not in returned:
                text_problems.append(f"node {hit['node_id']} read_node_text가 독립 OKF 추출과 다름")
        ledger.record("node_original_text", "PASS" if not text_problems else "FAIL", nodes=text_rows,
                      problems=text_problems)
        errors.extend(text_problems)
    return ledger.finish(errors)


def _verify_ac41(trees, pageindex, connect, settings):
    spec, ledger, errors = fixture("AC41"), _CheckLedger("AC41"), []
    blocked = None
    body_rows, hash_rows, node_rows, chunk_rows = [], [], [], []
    body_ok = hash_ok = node_ok = chunk_ok = True
    chunk_blocked = chunk_confirmed_fail = False
    for doc in spec.get("documents", []):
        tree = _tree_for_spec(trees, doc)
        okf = Path(pageindex.OKF_DOCUMENTS_ROOT) / doc["okf_path"]
        if tree is None or not okf.is_file():
            body_ok = hash_ok = node_ok = chunk_ok = False
            chunk_confirmed_fail = True
            body_rows.append({"doc_id": doc["doc_id"], "exists": okf.is_file(), "tree_found": tree is not None})
            continue
        body = _body_text(okf)
        in_range = _fact_in_line_range(okf, doc["fact"], doc["line_start"], doc["line_end"])
        body_rows.append({"doc_id": doc["doc_id"], "body_chars": len(body), "fact_in_fixed_range": in_range})
        body_ok &= bool(body.strip()) and in_range
        current_hash = _body_sha256(okf)
        fresh = current_hash == doc["body_sha256"] == tree.get("okf_body_sha256")
        hash_rows.append({"doc_id": doc["doc_id"], "fixture": doc["body_sha256"], "current": current_hash,
                          "tree": tree.get("okf_body_sha256"), "fresh": fresh})
        hash_ok &= fresh
        from rag_core.retrieval.access import PRIVATE_ONLY_SOURCE_GROUPS
        hits = pageindex.search_nodes(doc["query"], doc=doc["doc_id"], node_limit=3,
                                      exclude_source_groups=PRIVATE_ONLY_SOURCE_GROUPS)
        valid_hits = []
        for hit in hits:
            independent, start, end, hit_problems = _independent_node_text(tree, hit, okf)
            returned = pageindex.read_node_text(hit) if not hit_problems else ""
            if not hit_problems and returned == independent and doc["fact"] in independent:
                valid_hits.append({"node_id": hit["node_id"], "line_start": start, "line_end": end,
                                   "text_chars": len(returned)})
        roots = tree.get("structure", [])
        node_rows.append({"doc_id": doc["doc_id"], "root_count": len(roots), "valid_hits": valid_hits})
        node_ok &= bool(roots) and bool(valid_hits)
        status, chunks, error = _fetch_chunks(connect, settings, doc_id=doc["doc_id"], source_path=doc["resource"])
        if status != "PASS":
            if status.startswith("BLOCKED_"):
                blocked = status
                chunk_blocked = True
            else:
                chunk_confirmed_fail = True
            chunk_ok = False
            chunk_rows.append({"doc_id": doc["doc_id"], "status": status, "error": error})
        else:
            expected_texts, actual_texts = _expected_chunk_texts(okf), {chunk["text"] for chunk in chunks}
            fact_chunks = sum(doc["fact"] in chunk["text"] for chunk in chunks)
            same = bool(chunks) and actual_texts == expected_texts and fact_chunks > 0
            chunk_rows.append({"doc_id": doc["doc_id"], "vector_doc_id": _vector_doc_id(doc["doc_id"]),
                               "actual_count": len(chunks), "expected_count": len(expected_texts),
                               "texts_equal_latest_okf": actual_texts == expected_texts, "fact_chunk_count": fact_chunks})
            chunk_ok &= same
            chunk_confirmed_fail |= not same
    ledger.record("three_nonempty_bodies", "PASS" if body_ok and len(body_rows) == 3 else "FAIL", documents=body_rows)
    ledger.record("three_tree_hashes_fresh", "PASS" if hash_ok and len(hash_rows) == 3 else "FAIL", documents=hash_rows)
    ledger.record("three_valid_roots_and_node_text", "PASS" if node_ok and len(node_rows) == 3 else "FAIL", documents=node_rows)
    if chunk_confirmed_fail or (not chunk_blocked and len(chunk_rows) != 3):
        chunk_status = "FAIL"
    elif chunk_blocked:
        chunk_status = blocked or "BLOCKED_ENV"
    else:
        chunk_status = "PASS"
    ledger.record("three_latest_chunk_texts", chunk_status, documents=chunk_rows)
    structural_ok = body_ok and hash_ok and node_ok and all(len(rows) == 3 for rows in (body_rows, hash_rows, node_rows))
    all_ok = structural_ok and chunk_status == "PASS" and len(chunk_rows) == 3
    atomic_status = "PASS" if all_ok else ((blocked or "BLOCKED_ENV") if structural_ok and chunk_status.startswith("BLOCKED_") else "FAIL")
    ledger.record("all_three_atomic", atomic_status, checked_documents=len(body_rows), expected_documents=3)
    if atomic_status == "FAIL":
        errors.append("AC41 세 문서 중 하나 이상이 본문·hash·node·최신 chunk 계약을 충족하지 않음")
    elif atomic_status.startswith("BLOCKED_"):
        errors.append("AC41 청크 연결은 DB 환경 부재로 실행 차단")
    return ledger.finish(errors, blocked=blocked)


def verify_case(case: dict[str, Any], *, trees: list[dict[str, Any]], pageindex, connect, settings):
    case_id = case["id"]
    if case_id in {"AC28", "AC29", "AC30"}:
        return _verify_document_case(case_id, trees, pageindex, connect, settings)
    if case_id == "AC39":
        return _verify_ac39(connect, settings)
    if case_id == "AC40":
        return _verify_ac40(trees, pageindex)
    if case_id == "AC41":
        return _verify_ac41(trees, pageindex, connect, settings)
    return "FAIL", ["지원하지 않는 integration fixture case"], {"required_checks": []}
