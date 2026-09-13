# -*- coding: utf-8 -*-
"""v10.pptx 슬라이드16(0-idx, PPT표시 17번째) "글로벌수급지도 - 핵심광물지도"
(리튬 baseline)이 "수입" 라벨을 달고 있으면서 실제로는 "수출"(O) 모드
패턴("호주→중국" 1위)을 담고 있던 구버전(v9부터 이어진) 데이터를, 방금
komis.or.kr에서 직접 라이브 조회해 검증한 진짜 "수입"(I) 모드 데이터로
교체한다. 원본 스크립트(어느 세션이 이 슬라이드를 만들었는지)는 이제
추적 불가 — 정적 덤프(income_data/komis/komis_07_supply_map_global.json)
의 리튬 행이 전부 0건이라 거기서 온 것도 아니다."""
import json
import sys

sys.path.insert(0, "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad")
from pptx_helpers import find_pptx, get_shape_by_name, set_paragraph, clear_runs
from pptx import Presentation

DATA = json.load(open("/tmp/li_2026_map_global_verified.json", encoding="utf-8"))


def base_style(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            return run.font.size
    return None


def fill(shape, entries):
    size = base_style(shape)
    tf = shape.text_frame
    existing = list(tf.paragraphs)
    for i, (text, is_header) in enumerate(entries):
        p = existing[i] if i < len(existing) else tf.add_paragraph()
        set_paragraph(p, [(text, False)], size=size, bold=(True if is_header else None))
    for j in range(len(entries), len(existing)):
        clear_runs(existing[j])


def set_caption(shape, text):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = text
            for extra in para.runs[1:]:
                extra.text = ""
            return
    if shape.text_frame.paragraphs:
        set_paragraph(shape.text_frame.paragraphs[0], [(text, False)])


def set_table_value(table, row_label, value_text):
    for row in table.rows:
        if row.cells[0].text.strip() == row_label:
            cell = row.cells[1]
            for para in cell.text_frame.paragraphs:
                if para.runs:
                    para.runs[0].text = value_text
                    for extra in para.runs[1:]:
                        extra.text = ""
                    return True
    return False


path = find_pptx("v10")
p = Presentation(path)
slide = p.slides[16]
title = get_shape_by_name(slide, "제목 1").text_frame.text.strip()
assert title == "글로벌수급지도 - 핵심광물지도", title

core = DATA["summary"]["core_diagnosis"][0]["text"]
major = [s["text"] for s in DATA["summary"]["major_changes"]]
route_body = " ".join(major[:3])
korea_body = major[3] if len(major) > 3 else "대한민국이 포함된 교역 루트가 상위권에 없습니다."

box = get_shape_by_name(slide, "TextBox 4")
entries = [
    ("글로벌 교역 현황", True),
    (core + " ⚠2026년은 진행 중인 해라 KOMIS 통관 반영분이 아직 일부(연간 통상 규모 대비 약 10%대)만 집계돼 있어, 연말까지 순위·금액이 크게 달라질 수 있습니다.", False),
    ("주요 교역 루트", True),
    (route_body, False),
    ("한국 관련 루트", True),
    (korea_body, False),
]
fill(box, entries)

km = {m["id"]: m["value"] for m in DATA["key_metrics"]}
table = get_shape_by_name(slide, "표 6").table
set_table_value(table, "세계 교역 총액", "약 " + f"{km['total_amount']/1e4:,.2f}" + "만")
set_table_value(table, "1위 루트 비중", f"{km['top1_share_pct']:.2f}")
set_table_value(table, "상위3루트 비중", f"{km['top3_share_pct']:.2f}")
set_table_value(table, "상위5루트 비중", f"{km['top5_share_pct']:.2f}")

cap = get_shape_by_name(slide, "TextBox 10")
set_caption(
    cap,
    "검색 옵션 :광종 : 리튬, 수입, 2026, 수출입국가: 전체"
    " (★260913 재수정 — 구버전은 라벨은 '수입'이지만 실제 순위 패턴이"
    " '수출' 모드와 일치하는 구버전 데이터였음, komis.or.kr 실시간"
    " 재조회로 교체)",
)

p.save(path)
print("저장 완료:", path)
for shape in slide.shapes:
    if shape.has_text_frame and shape.text_frame.text.strip():
        print(shape.text_frame.text)
    if shape.has_table:
        for row in shape.table.rows:
            print(' ', row.cells[0].text.strip(), '|', row.cells[1].text.strip())
