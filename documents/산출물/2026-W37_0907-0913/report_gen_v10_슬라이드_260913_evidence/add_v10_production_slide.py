# -*- coding: utf-8 -*-
"""v10.pptx에 '광물지도 - 핵심광물지도 (생산량 검색)' 슬라이드를 새로
추가한다 — streamlit_demo의 "천톤" 기본값 버그를 고친 뒤, 그 수정된
경로로 실제 komis.or.kr 라이브 조회한 결과(measure=production)를 보여줘
버그가 실제로 해결됐음을 슬라이드로도 확인시킨다."""
import copy
import json
import sys

sys.path.insert(0, "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad")
from pptx_helpers import find_pptx, get_shape_by_name, set_paragraph, clear_runs
from pptx import Presentation

DATA = json.load(open(
    "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad/v10_map_mineral_production.json",
    encoding="utf-8",
))


def fmt(value, digits=2):
    if value is None:
        return ""
    if float(value).is_integer():
        return f"{int(value):,}"
    return f"{value:,.{digits}f}"


def pct(value, digits=2):
    return fmt(round(value, digits), digits)


def base_style(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            return run.font.size
    return None


def fill(shape, entries):
    size = base_style(shape)
    tf = shape.text_frame
    existing = list(tf.paragraphs)
    for i, (text, is_header) in enumerate(entries):
        p = existing[i] if i < len(existing) else tf.add_paragraph()
        set_paragraph(p, [(text, False)], size=size, bold=(True if is_header else None))
    for j in range(len(entries), len(existing)):
        clear_runs(existing[j])


def set_caption(shape, text):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = text
            for extra in para.runs[1:]:
                extra.text = ""
            return
    if shape.text_frame.paragraphs:
        set_paragraph(shape.text_frame.paragraphs[0], [(text, False)])


def set_title(slide, text):
    for shape in slide.shapes:
        if shape.name == "제목 1":
            tf = shape.text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = text
                for extra in tf.paragraphs[0].runs[1:]:
                    extra.text = ""
                for p in tf.paragraphs[1:]:
                    clear_runs(p)
                return
    raise KeyError("title shape not found")


def mid(metrics, metric_id):
    for m in metrics:
        if m["id"] == metric_id:
            return m
    raise KeyError(metric_id)


def set_table_value(table, row_label, value_text):
    for row in table.rows:
        if row.cells[0].text.strip() == row_label:
            cell = row.cells[1]
            for para in cell.text_frame.paragraphs:
                if para.runs:
                    para.runs[0].text = value_text
                    for extra in para.runs[1:]:
                        extra.text = ""
                    return True
    return False


def duplicate_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shape._element))
    return dest


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


path = find_pptx("v10")
p = Presentation(path)

# 슬라이드19(교차비교)를 템플릿으로 복제 — 같은 표/캡션 구조.
base_idx = 19
title19 = get_shape_by_name(p.slides[base_idx], "제목 1").text_frame.text.strip()
assert "교차비교" in title19, title19

new_slide = duplicate_slide(p, base_idx)
set_title(new_slide, "광물지도 - 핵심광물지도 (생산량 검색, 단위버그 수정 후 komis.or.kr 실시간조회)")

box = get_shape_by_name(new_slide, "TextBox 4")
major = DATA["sentences"]["major_changes"]
entries = [
    ("세계 생산량 현황", True), (" ".join(DATA["sentences"]["core_diagnosis"]), False),
    ("국가별 순위 및 변화(교차비교 포함)", True), (" ".join(major[:-1]), False),
    ("주요 변화", True), (major[-1], False),
]
fill(box, entries)

cap = get_shape_by_name(new_slide, "TextBox 10")
af = DATA["applied_filters"]
set_caption(
    cap,
    f"검색 옵션 :광종 : {af['mineral']}, 생산량, {af['start_year']}~{af['end_year']}, "
    "komis.or.kr 실시간조회(단위 자동유도, 천톤 기본값 버그 수정 후)",
)

table = get_shape_by_name(new_slide, "표 6").table
km = DATA["key_metrics"]
set_table_value(table, "현재 세계 매장량", "약 " + fmt(mid(km, "current_world_total")["value"] / 1e8) + "억")
set_table_value(table, "조회기간 세계합계 변화율", pct(mid(km, "period_world_total_change")["value"] * 100))
set_table_value(table, "1위 국가", mid(km, "top_country")["value"])
set_table_value(table, "1위 국가 비중", pct(mid(km, "top_country_share")["value"] * 100))
set_table_value(table, "상위 3개국 비중", pct(mid(km, "cr3")["value"] * 100))
set_table_value(table, "상위 5개국 비중", pct(mid(km, "cr5")["value"] * 100))
set_table_value(table, "1·2위 비중 차이", pct(mid(km, "top_two_share_gap")["value"] * 100))
set_table_value(table, "최대 감소 국가", mid(km, "max_decrease_country")["value"])
set_table_value(table, "전년 대비 변화율", pct(mid(km, "latest_year_total_change")["value"] * 100))
# 표 1행 라벨이 "현재 세계 매장량"으로 고정돼 있어(템플릿 원본, reserves
# 전제) 생산량 슬라이드에서는 혼동을 막기 위해 라벨 자체를 바꾼다.
for row in table.rows:
    if row.cells[0].text.strip() == "현재 세계 매장량":
        cell = row.cells[0]
        for para in cell.text_frame.paragraphs:
            if para.runs:
                para.runs[0].text = "현재 세계 생산량"
                for extra in para.runs[1:]:
                    extra.text = ""
                break
        break

move_slide(p, len(list(p.slides._sldIdLst)) - 1, base_idx + 1)

p.save(path)
print("저장 완료:", path, "| 총 슬라이드", len(p.slides))
