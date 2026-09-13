# -*- coding: utf-8 -*-
"""v10.pptx 슬라이드18(광물지도 baseline)·슬라이드19(매장량/생산량 교차비교)를
조회기간 2019~2025(7년, KOMIS 실제 제약을 넘음)에서 2021~2025(5년, 사용자
확인 — 광물지도 연간 검색은 5년까지만 가능)로 갱신한다. 두 슬라이드 모두
실제 report_gen 호출 결과(v10_map_mineral_2021_2025.json)로 교체."""
import json
import sys

sys.path.insert(0, "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad")
from pptx_helpers import find_pptx, get_shape_by_name, set_paragraph, clear_runs
from pptx import Presentation

DATA = json.load(open(
    "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad/v10_map_mineral_2021_2025.json",
    encoding="utf-8",
))


def fmt(value, digits=2):
    if value is None:
        return ""
    if float(value).is_integer():
        return f"{int(value):,}"
    return f"{value:,.{digits}f}"


def pct(value, digits=2):
    return fmt(round(value, digits), digits)


def base_style(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            return run.font.size
    return None


def fill(shape, entries):
    size = base_style(shape)
    tf = shape.text_frame
    existing = list(tf.paragraphs)
    for i, (text, is_header) in enumerate(entries):
        p = existing[i] if i < len(existing) else tf.add_paragraph()
        set_paragraph(p, [(text, False)], size=size, bold=(True if is_header else None))
    for j in range(len(entries), len(existing)):
        clear_runs(existing[j])


def set_caption(shape, text):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = text
            for extra in para.runs[1:]:
                extra.text = ""
            return
    if shape.text_frame.paragraphs:
        set_paragraph(shape.text_frame.paragraphs[0], [(text, False)])


def mid(metrics, metric_id):
    for m in metrics:
        if m["id"] == metric_id:
            return m
    raise KeyError(metric_id)


def set_table_value(table, row_label, value_text):
    for row in table.rows:
        if row.cells[0].text.strip() == row_label:
            cell = row.cells[1]
            for para in cell.text_frame.paragraphs:
                if para.runs:
                    para.runs[0].text = value_text
                    for extra in para.runs[1:]:
                        extra.text = ""
                    return True
    return False


def update_slide(slide, d, extra_caption=""):
    box = get_shape_by_name(slide, "TextBox 4")
    major = d["sentences"]["major_changes"]
    if extra_caption:  # 교차비교 슬라이드: major_changes 마지막에서 두번째가 cross_fact
        header2 = "국가별 순위 및 변화(교차비교 포함)"
        body2 = " ".join(major[:-1])
    else:
        header2 = "국가별 순위 및 변화"
        body2 = " ".join(major[:-1])
    entries = [
        ("세계 매장량 현황", True), (" ".join(d["sentences"]["core_diagnosis"]), False),
        (header2, True), (body2, False),
        ("주요 변화", True), (major[-1], False),
    ]
    fill(box, entries)

    cap = get_shape_by_name(slide, "TextBox 10")
    af = d["applied_filters"]
    set_caption(cap, f"검색 옵션 :광종 : {af['mineral']}, 매장량, {af['start_year']}~{af['end_year']}{extra_caption}")

    table = get_shape_by_name(slide, "표 6").table
    km = d["key_metrics"]
    set_table_value(table, "현재 세계 매장량", "약 " + fmt(mid(km, "current_world_total")["value"] / 1e8) + "억")
    set_table_value(table, "조회기간 세계합계 변화율", pct(mid(km, "period_world_total_change")["value"] * 100))
    set_table_value(table, "1위 국가", mid(km, "top_country")["value"])
    set_table_value(table, "1위 국가 비중", pct(mid(km, "top_country_share")["value"] * 100))
    set_table_value(table, "상위 3개국 비중", pct(mid(km, "cr3")["value"] * 100))
    set_table_value(table, "상위 5개국 비중", pct(mid(km, "cr5")["value"] * 100))
    set_table_value(table, "1·2위 비중 차이", pct(mid(km, "top_two_share_gap")["value"] * 100))
    set_table_value(table, "최대 감소 국가", mid(km, "max_decrease_country")["value"])
    set_table_value(table, "전년 대비 변화율", pct(mid(km, "latest_year_total_change")["value"] * 100))


path = find_pptx("v10")
p = Presentation(path)

slide18 = p.slides[18]
title18 = get_shape_by_name(slide18, "제목 1").text_frame.text.strip()
assert title18 == "광물지도 - 핵심광물지도", title18
update_slide(slide18, DATA["map_mineral_baseline_2021_2025"])

slide19 = p.slides[19]
title19 = get_shape_by_name(slide19, "제목 1").text_frame.text.strip()
assert "교차비교" in title19, title19
update_slide(slide19, DATA["map_mineral_cross_2021_2025"], extra_caption=", 교차비교: 생산량(komis_snapshot_response)")

p.save(path)
print("수정 완료:", path)
for idx in (18, 19):
    slide = p.slides[idx]
    print("#### slide", idx)
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            print(shape.text_frame.text)
