# -*- coding: utf-8 -*-
"""v10.pptx 슬라이드18·19·20(생산량 슬라이드는 index 20)을 report_gen
world-total 버그 수정(2026-09-13, 지도 국가합계 대신 KOMIS 공식 _TOTAL_ 사용)
반영한 재조회 결과로 갱신한다. 표 형식(수치 포맷)은 기존 스크립트와 동일하게
유지 — 값만 바뀐다."""
import json
import sys

sys.path.insert(0, "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad")
from pptx_helpers import find_pptx, get_shape_by_name, set_paragraph, clear_runs
from pptx import Presentation

DATA = json.load(open(
    "/tmp/claude-1002/-home-nuri-dev-git-ws-mine-ws-komir/859eef05-8b64-4e86-96f6-020f2796a86f/scratchpad/v10_map_mineral_2021_2025_worldtotal_fixed.json",
    encoding="utf-8",
))


def fmt(value, digits=2):
    if value is None:
        return ""
    if float(value).is_integer():
        return f"{int(value):,}"
    return f"{value:,.{digits}f}"


def pct(value, digits=2):
    return fmt(round(value, digits), digits)


def base_style(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            return run.font.size
    return None


def fill(shape, entries):
    size = base_style(shape)
    tf = shape.text_frame
    existing = list(tf.paragraphs)
    for i, (text, is_header) in enumerate(entries):
        p = existing[i] if i < len(existing) else tf.add_paragraph()
        set_paragraph(p, [(text, False)], size=size, bold=(True if is_header else None))
    for j in range(len(entries), len(existing)):
        clear_runs(existing[j])


def mid(metrics, metric_id):
    for m in metrics:
        if m["id"] == metric_id:
            return m
    raise KeyError(metric_id)


def set_table_value(table, row_label, value_text):
    for row in table.rows:
        if row.cells[0].text.strip() == row_label:
            cell = row.cells[1]
            for para in cell.text_frame.paragraphs:
                if para.runs:
                    para.runs[0].text = value_text
                    for extra in para.runs[1:]:
                        extra.text = ""
                    return True
    return False


def update_narrative(slide, d, cross=False):
    box = get_shape_by_name(slide, "TextBox 4")
    core = [s["text"] for s in d["summary"]["core_diagnosis"]]
    major = [s["text"] for s in d["summary"]["major_changes"]]
    header2 = "국가별 순위 및 변화(교차비교 포함)" if cross else "국가별 순위 및 변화"
    label = "세계 생산량 현황" if d["applied_filters"]["measure"] == "production" else "세계 매장량 현황"
    entries = [
        (label, True), (" ".join(core), False),
        (header2, True), (" ".join(major[:-1]), False),
        ("주요 변화", True), (major[-1], False),
    ]
    fill(box, entries)


def update_table(slide, d, total_row_label="현재 세계 매장량"):
    table = get_shape_by_name(slide, "표 6").table
    km = d["key_metrics"]
    set_table_value(table, total_row_label, "약 " + fmt(mid(km, "current_world_total")["value"] / 1e8) + "억")
    set_table_value(table, "조회기간 세계합계 변화율", pct(mid(km, "period_world_total_change")["value"] * 100))
    set_table_value(table, "1위 국가", mid(km, "top_country")["value"])
    set_table_value(table, "1위 국가 비중", pct(mid(km, "top_country_share")["value"] * 100))
    set_table_value(table, "상위 3개국 비중", pct(mid(km, "cr3")["value"] * 100))
    set_table_value(table, "상위 5개국 비중", pct(mid(km, "cr5")["value"] * 100))
    set_table_value(table, "1·2위 비중 차이", pct(mid(km, "top_two_share_gap")["value"] * 100))
    set_table_value(table, "최대 감소 국가", mid(km, "max_decrease_country")["value"])
    set_table_value(table, "전년 대비 변화율", pct(mid(km, "latest_year_total_change")["value"] * 100))


path = find_pptx("v10")
p = Presentation(path)

slide18 = p.slides[18]
assert get_shape_by_name(slide18, "제목 1").text_frame.text.strip() == "광물지도 - 핵심광물지도"
update_narrative(slide18, DATA["map_mineral_baseline_2021_2025"])
update_table(slide18, DATA["map_mineral_baseline_2021_2025"])

slide19 = p.slides[19]
assert "교차비교" in get_shape_by_name(slide19, "제목 1").text_frame.text.strip()
update_narrative(slide19, DATA["map_mineral_cross_2021_2025"], cross=True)
update_table(slide19, DATA["map_mineral_cross_2021_2025"])

slide20 = p.slides[20]
title20 = get_shape_by_name(slide20, "제목 1").text_frame.text.strip()
assert "생산량 검색" in title20, title20
update_narrative(slide20, DATA["map_mineral_production_2021_2025"])
update_table(slide20, DATA["map_mineral_production_2021_2025"], total_row_label="현재 세계 생산량")

p.save(path)
print("저장 완료:", path)
